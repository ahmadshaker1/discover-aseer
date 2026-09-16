#!/usr/bin/env python3
"""Create missing Directus *_cn fields (when absent) and fill them from
proofread Excel batches. Does not create or delete items. Unmatched sheet
rows are skipped. Hotel rows match existing English names only.

Usage:
  python3 scripts/update-cms-chinese-from-xlsx.py           # dry run
  python3 scripts/update-cms-chinese-from-xlsx.py --apply   # write to Directus
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from html import unescape
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
ENV_PATH = ROOT / ".env"

SHEETS = [
    Path(
        "/Users/ahmadshaker/Downloads/Aseer_Content_English_Chinese_2nd batch_proofreading completed.xlsx"
    ),
    Path(
        "/Users/ahmadshaker/Downloads/Aseer_Content_English_Chinese_3rd batch_proofreading completed.xlsx"
    ),
    Path(
        "/Users/ahmadshaker/Downloads/Chinese translation 4th Batch_Translated_proofreading completed and attention on hotel tab with possible duplication.xlsx"
    ),
    Path(
        "/Users/ahmadshaker/Downloads/Chinese Translation 5th Batch_Translated_proofreading completed.xlsx"
    ),
    Path(
        "/Users/ahmadshaker/Downloads/Chinese Translation 6th Batch_translated_proofreading completed.xlsx"
    ),
    Path(
        "/Users/ahmadshaker/Downloads/Chinese Translation 7th batch_translated_proofreading completed.xlsx"
    ),
]

# cn_field -> clone-from field (used only to copy type/interface)
CN_FIELD_SOURCES: dict[str, list[tuple[str, str]]] = {
    "seasons": [("title_cn", "title"), ("content_cn", "content")],
    "restaurants": [
        ("title_cn", "title_en"),
        ("content_cn", "content"),
        ("city_cn", "city"),
        ("cuisine_type_cn", "title_en"),
    ],
    "locations": [
        ("name_cn", "name_en"),
        ("category_cn", "category_en"),
        ("type_cn", "type_en"),
        ("description_cn", "description_en"),
    ],
    "experiences": [
        ("title_cn", "title_eng"),
        ("description_cn", "description_eng"),
        ("tour_agency_cn", "tour_agency_en"),
        ("duration_cn", "duration_En"),
        ("type_cn", "type_en"),
    ],
    "events": [
        ("title_cn", "title_en"),
        ("description_cn", "description_en"),
        ("city_cn", "city_en"),
        ("type_cn", "type_en"),
    ],
    "accomodation": [
        ("name_cn", "name_en"),
        ("content_cn", "content"),
        ("city_cn", "city_en"),
        ("type_cn", "type"),
    ],
    "cuisine": [
        ("title_cn", "title_en"),
        ("hero_content_cn", "hero_content_en"),
        ("subtitle_cn", "subtitle_en"),
        ("subtitle_purple_cn", "subtitle_purple_en"),
        ("content_cn", "content_en"),
        ("extra_content_cn", "extra_content_en"),
        ("cuisine_type_cn", "title_en"),
    ],
    "films": [("title_cn", "title_en"), ("type_cn", "type")],
    "tourism_providers": [("title_cn", "title_en")],
    "support_service": [
        ("title_cn", "title_en"),
        ("city_cn", "city_en"),
        ("type_cn", "type"),
    ],
    "privacy_policy": [("privacy_policy_cn", "privacy_policy")],
    "attractions": [],  # already present
    "destination": [],
    "pages": [],
}

CITY_ALIASES = {
    "ahad rafida": "ahad rafidah",
    "dhahran al janub": "dhahran al janoub",
    "al haridhah": "al haridah",
    "almajaridah": "al majardah",
    "al majaridah": "al majardah",
    "bariq": "bareq",
    "namas": "al namas",
}

ARRAY_OR_JSON_TYPES = {"json", "csv", "alias", "o2m", "m2o", "m2m", "files", "file"}


def load_env() -> None:
    if not ENV_PATH.exists():
        return
    for line in ENV_PATH.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def cell(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"none", "nan", "null"}:
        return ""
    return text


def norm(value: str) -> str:
    text = unescape(cell(value))
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.lower()
    text = re.sub(r"[\u200f\u200e]", "", text)
    text = re.sub(r"[^\w\u4e00-\u9fff]+", " ", text, flags=re.UNICODE)
    return re.sub(r"\s+", " ", text).strip()


def city_key(value: str) -> str:
    key = norm(value)
    return CITY_ALIASES.get(key, key)


def merge_patch(dest: dict, incoming: dict) -> None:
    for key, value in incoming.items():
        if cell(value):
            dest[key] = cell(value)


def request(method: str, url: str, token: str | None = None, body=None, timeout=90):
    data = None
    headers = {}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    if body is not None:
        data = json.dumps(body).encode("utf-8")
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            raw = resp.read().decode("utf-8")
            return resp.status, json.loads(raw) if raw.strip() else {}
    except urllib.error.HTTPError as exc:
        raw = exc.read().decode("utf-8", errors="replace")
        try:
            parsed = json.loads(raw) if raw.strip() else {}
        except json.JSONDecodeError:
            parsed = {"raw": raw[:300]}
        return exc.code, parsed


def fetch_all_items(
    base: str, token: str | None, collection: str, fields: str
) -> list[dict]:
    items: list[dict] = []
    page = 1
    while True:
        qs = urllib.parse.urlencode({"fields": fields, "limit": 100, "page": page})
        status, payload = request("GET", f"{base}/items/{collection}?{qs}", token)
        if status != 200:
            raise RuntimeError(f"GET {collection} failed ({status})")
        data = payload.get("data")
        if isinstance(data, dict):
            items.append(data)
            break
        batch = data or []
        items.extend(batch)
        if len(batch) < 100:
            break
        page += 1
        if page > 50:
            break
    return items


def load_workbook_rows(path: Path) -> dict[str, list[list]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: dict[str, list[list]] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
        out[name] = rows
    wb.close()
    return out


def list_fields(base: str, token: str, collection: str) -> list[dict]:
    status, payload = request("GET", f"{base}/fields/{collection}", token)
    if status != 200:
        print(f"! could not list fields for {collection}: {status}")
        return []
    return payload.get("data") or []


def ensure_cn_fields(base: str, token: str, apply: bool) -> dict[str, set[str]]:
    """Create missing *_cn columns. Returns cn field names per collection."""
    present: dict[str, set[str]] = {}
    for collection, specs in CN_FIELD_SOURCES.items():
        rows = list_fields(base, token, collection)
        by_name = {row.get("field"): row for row in rows if row.get("field")}
        cn_names = {name for name in by_name if str(name).endswith("_cn")}
        for cn_field, source_field in specs:
            if cn_field in by_name:
                cn_names.add(cn_field)
                continue
            source = by_name.get(source_field) or {}
            field_type = source.get("type") or "text"
            if field_type in ARRAY_OR_JSON_TYPES:
                field_type = "text"
            interface = (source.get("meta") or {}).get("interface") or "input"
            if interface in {
                "select-dropdown-m2o",
                "list-m2m",
                "file",
                "files",
                "select-dropdown",
                "select-multiple-dropdown",
            }:
                interface = "input"
            print(f"+ schema {collection}.{cn_field} (from {source_field}: {field_type}/{interface})")
            if not apply:
                continue
            cn_names.add(cn_field)
            body = {
                "field": cn_field,
                "type": field_type,
                "schema": {"is_nullable": True},
                "meta": {
                    "interface": interface,
                    "hidden": False,
                    "width": "full",
                    "note": "Chinese translation",
                },
            }
            st, res = request("POST", f"{base}/fields/{collection}", token, body)
            if st not in (200, 201, 204):
                print(f"  ! failed to create {collection}.{cn_field}: {st} {res}")
                cn_names.discard(cn_field)
            time.sleep(0.2)
        present[collection] = cn_names
    return present


def collect_translations() -> dict[str, dict[str, dict[str, str]]]:
    buckets: dict[str, dict[str, dict[str, str]]] = {
        name: {}
        for name in [
            "seasons",
            "restaurants",
            "locations",
            "experiences",
            "events",
            "accomodation",
            "cuisine",
            "films",
            "tourism_providers",
            "support_service",
            "pages",
            "faq",
            "privacy_policy",
            "attractions",
        ]
    }

    def put(collection: str, english: str, patch: dict[str, str]) -> None:
        key = norm(english)
        cleaned = {k: cell(v) for k, v in patch.items() if cell(v)}
        if not key or not cleaned:
            return
        merge_patch(buckets[collection].setdefault(key, {}), cleaned)

    for path in SHEETS:
        if not path.exists():
            print(f"! missing workbook {path.name}")
            continue
        sheets = load_workbook_rows(path)
        name = path.name.lower()

        if "2nd batch" in name:
            for row in sheets.get("Aseeri Cusine", [])[1:]:
                put(
                    "cuisine",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "title_cn": row[1] if len(row) > 1 else "",
                        "hero_content_cn": row[3] if len(row) > 3 else "",
                        "subtitle_cn": row[5] if len(row) > 5 else "",
                        "subtitle_purple_cn": row[7] if len(row) > 7 else "",
                        "content_cn": row[9] if len(row) > 9 else "",
                        "extra_content_cn": row[11] if len(row) > 11 else "",
                        "cuisine_type_cn": row[13] if len(row) > 13 else "",
                    },
                )
            for row in sheets.get("Aseer Film", [])[1:]:
                put(
                    "films",
                    cell(row[2] if len(row) > 2 else None),
                    {
                        "title_cn": row[3] if len(row) > 3 else "",
                        "type_cn": row[1] if len(row) > 1 else "",
                    },
                )
            for row in sheets.get("Tourism Companies", [])[1:]:
                put(
                    "tourism_providers",
                    cell(row[0] if len(row) > 0 else None),
                    {"title_cn": row[1] if len(row) > 1 else ""},
                )
            for row in sheets.get("Support Services", [])[1:]:
                put(
                    "support_service",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "title_cn": row[1] if len(row) > 1 else "",
                        "type_cn": row[5] if len(row) > 5 else "",
                        "city_cn": row[7] if len(row) > 7 else "",
                    },
                )

        if "3rd batch" in name:
            for row in sheets.get("Seasons", [])[1:]:
                put(
                    "seasons",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "title_cn": row[1] if len(row) > 1 else "",
                        "content_cn": row[3] if len(row) > 3 else "",
                    },
                )
            for row in sheets.get("Restaurants", [])[1:]:
                put(
                    "restaurants",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "title_cn": row[1] if len(row) > 1 else "",
                        "content_cn": row[3] if len(row) > 3 else "",
                        "cuisine_type_cn": row[5] if len(row) > 5 else "",
                        "city_cn": row[7] if len(row) > 7 else "",
                    },
                )
            for row in sheets.get("Locations", [])[1:]:
                put(
                    "locations",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "name_cn": row[1] if len(row) > 1 else "",
                        "category_cn": row[3] if len(row) > 3 else "",
                        "type_cn": row[5] if len(row) > 5 else "",
                        "description_cn": row[7] if len(row) > 7 else "",
                    },
                )
                put(
                    "attractions",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "name_cn": row[1] if len(row) > 1 else "",
                        "type_cn": row[5] if len(row) > 5 else "",
                        "content_cn": row[7] if len(row) > 7 else "",
                        "content_home_page_card_content_cn": row[7] if len(row) > 7 else "",
                    },
                )
            for row in sheets.get("Experiences", [])[1:]:
                put(
                    "experiences",
                    cell(row[4] if len(row) > 4 else None),
                    {
                        "tour_agency_cn": row[1] if len(row) > 1 else "",
                        "type_cn": row[3] if len(row) > 3 else "",
                        "title_cn": row[5] if len(row) > 5 else "",
                        "description_cn": row[7] if len(row) > 7 else "",
                        "duration_cn": row[9] if len(row) > 9 else "",
                    },
                )
            for row in sheets.get("Events", [])[1:]:
                put(
                    "events",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "title_cn": row[1] if len(row) > 1 else "",
                        "description_cn": row[3] if len(row) > 3 else "",
                        "type_cn": row[5] if len(row) > 5 else "",
                        "city_cn": row[7] if len(row) > 7 else "",
                    },
                )
            for row in sheets.get("Acommodations", [])[1:]:
                put(
                    "accomodation",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "name_cn": row[1] if len(row) > 1 else "",
                        "content_cn": row[3] if len(row) > 3 else "",
                        "city_cn": row[5] if len(row) > 5 else "",
                        "type_cn": row[7] if len(row) > 7 else "",
                    },
                )

        if "4th batch" in name:
            for row in sheets.get("Events", [])[1:]:
                put(
                    "events",
                    cell(row[1] if len(row) > 1 else None),
                    {
                        "title_cn": row[0] if len(row) > 0 else "",
                        "description_cn": row[2] if len(row) > 2 else "",
                        "type_cn": row[5] if len(row) > 5 else "",
                        "city_cn": row[7] if len(row) > 7 else "",
                    },
                )
            for row in sheets.get("Hotels", [])[1:]:
                put(
                    "accomodation",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "name_cn": row[1] if len(row) > 1 else "",
                        "content_cn": row[3] if len(row) > 3 else "",
                        "city_cn": row[7] if len(row) > 7 else "",
                    },
                )

        if "5th batch" in name:
            for row in sheets.get("Tourism Companies", [])[1:]:
                put(
                    "tourism_providers",
                    cell(row[0] if len(row) > 0 else None),
                    {"title_cn": row[1] if len(row) > 1 else ""},
                )
            for row in sheets.get("FAQ", [])[1:]:
                put(
                    "faq",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "question_cn": row[1] if len(row) > 1 else "",
                        "answer_cn": row[3] if len(row) > 3 else "",
                    },
                )

        if "6th batch" in name:
            bodies = []
            for row in sheets.get("Privacy Policy", [])[1:]:
                title = cell(row[1] if len(row) > 1 else None)
                body = cell(row[3] if len(row) > 3 else None)
                if title or body:
                    bodies.append(f"<h2>{title}</h2>\n{body}" if title else body)
            if bodies:
                put(
                    "privacy_policy",
                    "introduction",
                    {"privacy_policy_cn": "\n".join(bodies)},
                )

        if "7th batch" in name:
            for row in sheets.get("Pages", [])[1:]:
                put(
                    "pages",
                    cell(row[0] if len(row) > 0 else None),
                    {
                        "title_cn": row[1] if len(row) > 1 else "",
                        "subtitle_cn": row[3] if len(row) > 3 else "",
                    },
                )

    return buckets


def needed_patch(item: dict, patch: dict[str, str]) -> dict[str, str]:
    out = {}
    for key, value in patch.items():
        if cell(value) and cell(item.get(key)) != cell(value):
            out[key] = cell(value)
    return out


def patch_item(base, token, collection, item_id, patch, apply) -> bool:
    if not patch:
        return False
    if not apply:
        return True
    st, _res = request(
        "PATCH",
        f"{base}/items/{collection}/{urllib.parse.quote(str(item_id), safe='')}",
        token,
        patch,
    )
    if st not in (200, 204):
        print(f"  ! PATCH {collection}/{item_id} failed ({st})")
        return False
    time.sleep(0.05)
    return True


def update_simple(
    base,
    read_token,
    write_token,
    collection,
    english_keys,
    translations,
    cn_fields,
    apply,
) -> tuple[int, int, int]:
    if not translations:
        return (0, 0, 0)
    get_fields = ["id", *english_keys, *sorted(cn_fields)]
    fields = ",".join(get_fields)
    try:
        items = fetch_all_items(base, read_token, collection, fields)
    except RuntimeError:
        try:
            items = fetch_all_items(base, write_token, collection, fields)
        except RuntimeError as exc:
            print(f"  {collection}: skip ({exc})")
            return (0, 0, 0)
    index: dict[str, list[dict]] = {}
    for item in items:
        for key in english_keys:
            n = norm(item.get(key) or "")
            if n:
                index.setdefault(n, []).append(item)

    updated = unmatched = skipped = 0
    used_ids = set()
    for english_key, patch in translations.items():
        matches = index.get(english_key, [])
        if not matches:
            unmatched += 1
            continue
        allowed = {
            k: v
            for k, v in patch.items()
            if cell(v) and (not cn_fields or k in cn_fields)
        }
        if not allowed:
            skipped += 1
            continue
        for item in matches:
            ident = item.get("id")
            if ident in used_ids:
                continue
            used_ids.add(ident)
            current = needed_patch(item, allowed)
            if not current:
                skipped += 1
                continue
            if patch_item(base, write_token, collection, ident, current, apply):
                updated += 1
            else:
                skipped += 1
    return updated, unmatched, skipped


def update_faq(base, token, translations, apply) -> tuple[int, int]:
    if not translations:
        return 0, 0
    try:
        items = fetch_all_items(base, token, "faq", "id,questions")
    except RuntimeError:
        print("! faq not readable")
        return 0, 0
    updated_questions = 0
    unmatched = 0
    for item in items:
        questions = item.get("questions")
        if not isinstance(questions, list):
            continue
        changed = False
        for q in questions:
            if not isinstance(q, dict):
                continue
            key = norm(q.get("question_en") or "")
            patch = translations.get(key)
            if not patch:
                unmatched += 1
                continue
            if patch.get("question_cn") and q.get("question_cn") != patch["question_cn"]:
                q["question_cn"] = patch["question_cn"]
                changed = True
                updated_questions += 1
            if patch.get("answer_cn") and q.get("answer_cn") != patch["answer_cn"]:
                q["answer_cn"] = patch["answer_cn"]
                changed = True
        if changed:
            ok = patch_item(
                base, token, "faq", item.get("id"), {"questions": questions}, apply
            )
            if not ok and apply:
                st, _res = request("PATCH", f"{base}/items/faq", token, {"questions": questions})
                if st not in (200, 204):
                    print(f"  ! PATCH faq singleton failed ({st})")
                else:
                    time.sleep(0.05)
    return updated_questions, unmatched


def update_privacy(base, token, translations, apply) -> int:
    html = next(iter(translations.values()), {}).get("privacy_policy_cn")
    if not html:
        return 0
    try:
        items = fetch_all_items(base, token, "privacy_policy", "id,privacy_policy_cn")
    except RuntimeError:
        print("! privacy_policy has no readable items; skip (no create)")
        return 0
    if not items:
        print("! privacy_policy has no existing items; skip (no create)")
        return 0
    count = 0
    for item in items:
        if patch_item(
            base,
            token,
            "privacy_policy",
            item.get("id"),
            {"privacy_policy_cn": html},
            apply,
        ):
            count += 1
    return count


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    load_env()
    base = (
        os.environ.get("DIRECTUS_WRITE_BASE_URL")
        or os.environ.get("NEXT_PUBLIC_DIRECTUS_APP_URL")
        or "https://tool-portal.discoveraseer.com"
    ).rstrip("/")
    token = os.environ.get("DIRECTUS_ADMIN_TOKEN") or ""
    if args.apply and not token:
        print("Missing Directus admin token")
        return 1

    print("Loading workbooks…")
    translations = collect_translations()
    for col, rows in translations.items():
        print(f"  sheet keys {col}: {len(rows)}")

    print("Ensuring _cn fields…", "(dry run)" if not args.apply else "(APPLY)")
    cn_by_collection: dict[str, set[str]] = {}
    if token:
        cn_by_collection = ensure_cn_fields(base, token, args.apply)
    else:
        print("  no token; assuming fields will be created on --apply")

    print("Updating items…", "(APPLY)" if args.apply else "(dry run)")
    plans = [
        ("seasons", ["title", "title_ar"]),
        ("restaurants", ["title_en", "title_ar"]),
        ("locations", ["name_en", "name_ar"]),
        ("experiences", ["title_eng", "title"]),
        ("events", ["title_en", "title"]),
        ("accomodation", ["name_en", "name_ar"]),
        ("cuisine", ["title_en", "title_ar"]),
        ("films", ["title_en", "title_ar"]),
        ("tourism_providers", ["title_en", "title_ar"]),
        ("support_service", ["title_en", "title_ar"]),
        ("attractions", ["name_en"]),
        ("pages", ["title_en", "title"]),
    ]
    read_token = None
    write_token = token or None
    for collection, keys in plans:
        updated, unmatched, skipped = update_simple(
            base,
            read_token,
            write_token,
            collection,
            keys,
            translations.get(collection, {}),
            cn_by_collection.get(collection, set()),
            args.apply,
        )
        print(
            f"  {collection}: updated={updated} unmatched_sheet_rows={unmatched} skipped={skipped}"
        )

    faq_updated, faq_unmatched = update_faq(
        base, write_token, translations.get("faq", {}), args.apply
    )
    print(f"  faq: questions_updated={faq_updated} unmatched={faq_unmatched}")
    privacy_updated = update_privacy(
        base, write_token, translations.get("privacy_policy", {}), args.apply
    )
    print(f"  privacy_policy: updated={privacy_updated}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
